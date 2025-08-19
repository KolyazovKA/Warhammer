import os
import docx
import pdfplumber
# import markdown
# import openpyxl
# import xlrd

def load_file_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()

    # if ext == ".docx":
    #     doc = docx.Document(file_path)
    #     return "\n".join([p.text for p in doc.paragraphs])

    # elif ext == ".doc":
    #     # обрабатываем через xlrd-like для старого doc? проще: not supported напрямую
    #     raise ValueError(".doc не поддерживается напрямую, конвертируйте в .docx")

    if ext == ".pdf":
        text = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text.append(page.extract_text() or "")
        return "\n".join(text)

    # elif ext == ".md":
    #     with open(file_path, "r", encoding="utf-8") as f:
    #         raw = f.read()
    #     return markdown.markdown(raw)

    # # elif ext in [".xls", ".xlsx"]:
    # #     text = []
    # #     if ext == ".xlsx":
    # #         wb = openpyxl.load_workbook(file_path)
    # #         for sheet in wb:
    # #             for row in sheet.iter_rows(values_only=True):
    # #                 text.append(" ".join([str(c) for c in row if c]))
    # #     else:
    # #         wb = xlrd.open_workbook(file_path)
    # #         for sheet in wb.sheets():
    # #             for row_idx in range(sheet.nrows):
    # #                 row = sheet.row_values(row_idx)
    # #                 text.append(" ".join([str(c) for c in row if c]))
    # #     return "\n".join(text)

    else:
        raise ValueError(f"Формат {ext} не поддерживается")
